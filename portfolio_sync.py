#!/usr/bin/env python3
"""
Sincronizzazione posizioni aperte dal dossier titoli (export XLS del broker).

Legge il dossier (formato "Movimenti Dossier Titoli": Operazione, Data
valuta, Descrizione, Titolo, ISIN, Segno A/V, Quantità, Divisa, Prezzo,
Cambio, Controvalore), ricostruisce le posizioni aperte in FIFO per ISIN e
le mappa a ticker TradingView.

Il repo e' pubblico: dossier XLS e mappa ISIN sono dati personali, quindi
non vengono committati in chiaro. Fonte primaria = GitHub Secret
(MOVIMENTI_XLSX_B64 in base64, ISIN_MAP_JSON come stringa JSON), mai
scritto su disco nel repo. Se il secret non c'e' (uso locale/di test),
si legge il file committato in data/ — che quindi deve restare un
template vuoto/pseudonimo, non i tuoi dati reali.

Le posizioni chiuse (quantita residua zero dopo il FIFO) non vengono
restituite: spariscono dal report senza bisogno di un intervento manuale.
"""
import base64
import io
import json
import logging
import os
from collections import deque
from datetime import date, datetime

logger = logging.getLogger(__name__)

MOVIMENTI_PATH = os.environ.get("MOVIMENTI_PATH", "data/movimenti_dossier.xlsx")
ISIN_MAP_PATH = os.environ.get("ISIN_MAP_PATH", "data/isin_map.json")
MOVIMENTI_XLSX_B64 = os.environ.get("MOVIMENTI_XLSX_B64", "")
ISIN_MAP_JSON = os.environ.get("ISIN_MAP_JSON", "")


def dossier_disponibile():
    """True se c'e' un dossier da sincronizzare, come secret o come file
    committato. Usato da generate_report.py per decidere se sincronizzare
    dal dossier o ricadere sul vecchio secret PORTFOLIO_JSON."""
    return bool(MOVIMENTI_XLSX_B64) or os.path.exists(MOVIMENTI_PATH)

# Intestazioni come compaiono nell'export "Movimenti Dossier Titoli", normalizzate
# in minuscolo. L'ordine delle colonne nel file non conta: si legge per nome.
_HEADER_ALIASES = {
    "operazione": "operazione",
    "data valuta": "data_valuta",
    "descrizione": "descrizione",
    "titolo": "titolo",
    "isin": "isin",
    "segno a/v": "segno",
    "quantità": "quantita",
    "quantita": "quantita",
    "divisa": "divisa",
    "prezzo": "prezzo",
    "cambio": "cambio",
    "controvalore": "controvalore",
}


def _norm_header(h):
    return _HEADER_ALIASES.get(str(h or "").strip().lower())


def _parse_number(v):
    """Coerce numeri che arrivano come stringa con virgola decimale (export italiani)."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            continue
    return None


def read_movimenti(path=None):
    """Legge il dossier XLS (secret MOVIMENTI_XLSX_B64 se presente, altrimenti
    il file committato) e restituisce una lista di movimenti normalizzati.
    Righe non riconducibili a un acquisto/vendita di un titolo (cedole,
    commissioni, bonifici...) vengono scartate silenziosamente: non hanno
    un ISIN, un Segno A/V valido o una quantita'."""
    path = path or MOVIMENTI_PATH
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.error("openpyxl non installato: impossibile leggere il dossier XLS")
        return []

    if MOVIMENTI_XLSX_B64:
        try:
            sorgente = io.BytesIO(base64.b64decode(MOVIMENTI_XLSX_B64))
            logger.info("Dossier movimenti letto dal secret MOVIMENTI_XLSX_B64")
        except Exception as e:
            logger.error(f"MOVIMENTI_XLSX_B64 illeggibile ({e})")
            return []
    elif os.path.exists(path):
        sorgente = path
    else:
        logger.warning(f"Dossier movimenti non trovato: né secret MOVIMENTI_XLSX_B64 né {path}")
        return []

    wb = load_workbook(sorgente, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    cols = [_norm_header(h) for h in header]

    out = []
    for raw in rows:
        rec = {c: v for c, v in zip(cols, raw) if c}  # scarta colonne non riconosciute
        isin = str(rec.get("isin") or "").strip()
        segno = str(rec.get("segno") or "").strip().upper()
        qty = _parse_number(rec.get("quantita"))
        if not isin or segno not in ("A", "V") or not qty:
            continue
        data_val = _parse_date(rec.get("data_valuta"))
        prezzo = _parse_number(rec.get("prezzo"))
        if not data_val or prezzo is None:
            continue
        out.append({
            "isin": isin,
            "titolo": str(rec.get("titolo") or rec.get("descrizione") or isin).strip(),
            "segno": segno,
            "quantita": abs(qty),
            "prezzo": prezzo,
            "divisa": str(rec.get("divisa") or "EUR").strip(),
            "data_valuta": data_val,
        })
    logger.info(f"Movimenti letti: {len(out)} righe valide (acquisto/vendita titoli)")
    return out


def build_fifo_positions(movimenti):
    """Ricostruisce le posizioni aperte in FIFO, per ISIN.
    Le vendite consumano i lotti piu' vecchi; le posizioni chiuse (quantita
    residua zero) non compaiono nel risultato."""
    by_isin = {}
    for m in movimenti:
        by_isin.setdefault(m["isin"], []).append(m)

    positions = []
    for isin, righe in by_isin.items():
        righe = sorted(righe, key=lambda r: r["data_valuta"])
        lots = deque()  # ognuno: {"qty", "prezzo", "data"}
        titolo, divisa = righe[-1]["titolo"], righe[-1]["divisa"]

        for r in righe:
            if r["segno"] == "A":
                lots.append({"qty": r["quantita"], "prezzo": r["prezzo"], "data": r["data_valuta"]})
            else:  # "V"
                da_vendere = r["quantita"]
                while da_vendere > 1e-9 and lots:
                    lotto = lots[0]
                    consumato = min(lotto["qty"], da_vendere)
                    lotto["qty"] -= consumato
                    da_vendere -= consumato
                    if lotto["qty"] <= 1e-9:
                        lots.popleft()
                if da_vendere > 1e-9:
                    # Vendita superiore ai lotti registrati: probabile acquisto
                    # precedente all'inizio dello storico esportato. Segnala e
                    # ignora l'eccedenza, senza far esplodere il calcolo.
                    logger.warning(
                        f"{isin}: vendute {da_vendere:.2f} unita' senza lotto "
                        f"corrispondente (storico movimenti incompleto?)"
                    )

        if not lots:
            continue  # posizione chiusa: non compare nel report

        tot_qty = sum(l["qty"] for l in lots)
        avg_price = sum(l["qty"] * l["prezzo"] for l in lots) / tot_qty
        # Data di apertura media, pesata per quantita' sui lotti ancora aperti
        avg_ord = sum(l["qty"] * l["data"].toordinal() for l in lots) / tot_qty
        entry_date = date.fromordinal(round(avg_ord))
        holding_days = (date.today() - entry_date).days

        positions.append({
            "isin": isin,
            "titolo": titolo,
            "divisa": divisa,
            "quantita": round(tot_qty, 4),
            "prezzo_medio": round(avg_price, 4),
            "data_apertura": entry_date.isoformat(),
            "giorni_detenzione": holding_days,
        })

    positions.sort(key=lambda p: p["giorni_detenzione"], reverse=True)
    logger.info(f"Posizioni aperte ricostruite: {len(positions)}")
    return positions


def load_isin_map(path=None):
    """Mappa ISIN -> ticker TradingView: secret ISIN_MAP_JSON se presente,
    altrimenti il file committato (che deve restare un template vuoto)."""
    if ISIN_MAP_JSON:
        try:
            return json.loads(ISIN_MAP_JSON)
        except Exception as e:
            logger.warning(f"ISIN_MAP_JSON illeggibile ({e})")
            return {}
    path = path or ISIN_MAP_PATH
    if not os.path.exists(path):
        logger.warning(f"Mappa ISIN→TradingView non trovata: né secret ISIN_MAP_JSON né {path}")
        return {}
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        logger.warning(f"Mappa ISIN→TradingView illeggibile ({e})")
        return {}


def sync_portfolio():
    """Punto d'ingresso: dossier XLS → posizioni aperte → ticker TradingView.
    Restituisce la stessa forma di lista che il resto della pipeline si aspetta
    da PORTFOLIO ([{"symbol","name","type",...}]), con i campi FIFO aggiunti."""
    movimenti = read_movimenti()
    posizioni = build_fifo_positions(movimenti)
    isin_map = load_isin_map()

    non_mappati = []
    out = []
    for p in posizioni:
        mapped = isin_map.get(p["isin"])
        if not mapped:
            non_mappati.append(f'{p["isin"]} ({p["titolo"]})')
        out.append({
            "symbol": (mapped or {}).get("tv_symbol", p["isin"]),
            "name": (mapped or {}).get("name", p["titolo"]),
            "type": (mapped or {}).get("type", "Azione"),
            "isin": p["isin"],
            "quantita": p["quantita"],
            "prezzo_medio": p["prezzo_medio"],
            "divisa": p["divisa"],
            "data_apertura": p["data_apertura"],
            "giorni_detenzione": p["giorni_detenzione"],
        })
    if non_mappati:
        logger.warning(
            f"ISIN senza ticker TradingView in {ISIN_MAP_PATH} (dati tecnici "
            "assenti per questi titoli finche' non li mappi): " + "; ".join(non_mappati)
        )
    return out
